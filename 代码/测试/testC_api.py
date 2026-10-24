# -*- coding: utf-8 -*-
"""全套测试 · 第三批：HTTP 接口层 + CLI 命令行。"""
import json
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _testenv as ENV  # noqa: E402

PROJ = ENV.activate()

# ── 自带服务：用当前代码起一个，跑完自动关（避免连到残留旧进程）──
import server as _server  # noqa: E402

_server_handle = _server.ensure(force_own=False)
BASE = _server.BASE
B = BASE
import urllib.error
import urllib.parse
import urllib.request
import uuid

R = {"pass": [], "fail": []}


def ck(name, ok, detail=""):
    (R["pass"] if ok else R["fail"]).append({"name": name, "detail": str(detail)[:200]})
    print("  %s %s%s" % ("✅" if ok else "❌", name, ("　→ " + str(detail)[:90]) if detail else ""))


def get(p, timeout=120):
    with urllib.request.urlopen(B + p, timeout=timeout) as r:
        return r.status, json.load(r)


def get_headers(path, timeout=60):
    """返回小写键的响应头（HTTP 头名大小写不敏感）。"""
    with urllib.request.urlopen(B + path, timeout=timeout) as r:
        return r.status, {k.lower(): v for k, v in r.headers.items()}


def post_noredirect(p, fields=None, timeout=120):
    """POST 但不跟随重定向，用于验证 303。"""
    class NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *a, **k):
            return None
    op = urllib.request.build_opener(NoRedirect)
    data = urllib.parse.urlencode(fields or {}).encode()
    try:
        with op.open(B + p, data=data, timeout=timeout) as r:
            return r.status, dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers)


def post(p, fields=None, timeout=300):
    data = urllib.parse.urlencode(fields or {}).encode()
    with urllib.request.urlopen(urllib.request.Request(B + p, data=data, method="POST"),
                                timeout=timeout) as r:
        return r.status, json.load(r)


def post_raw(p, body, ctype, timeout=300):
    with urllib.request.urlopen(urllib.request.Request(B + p, data=body,
                                headers={"Content-Type": ctype}, method="POST"),
                                timeout=timeout) as r:
        return r.status, json.load(r)


def expect_http(p, method="GET", fields=None, code=None):
    """期望特定 HTTP 状态码；返回 (是否符合期望, 实际状态码)。默认期望 200。"""
    want = code if code is not None else 200
    try:
        data = urllib.parse.urlencode(fields or {}).encode() if method == "POST" else None
        req = urllib.request.Request(B + p, data=data, method=method)
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status == want, r.status
    except urllib.error.HTTPError as e:
        return e.code == want, e.code
    except Exception as e:
        return False, repr(e)[:60]


# ══════════ 1. 只读接口 ══════════
print("=" * 70)
print("1. 只读接口：state / papers / privacy / missing / evidence")
print("=" * 70)
st, j = get("/api/state")
ck("GET /api/state 返回 200", st == 200)
for k in ("settings", "cost", "freshness", "corpus", "docs", "cards", "matrix",
          "draft", "check", "check_issues", "benchmark", "has_docx"):
    ck("state 含字段 %s" % k, k in j)
ck("state.drafts 有参考文献", len(j["draft"]["refs"]) > 0, len(j["draft"]["refs"]))
# 先让产物与文献库对齐，再断言"新鲜度判定为最新"。
# 原来直接断言，等于在测"上一个脚本有没有刷过稿子"；一旦数据是刚还原的
# 快照（draft 落后于库），这里就会红，掩盖真正的问题。
if j["freshness"].get("needs_run") is not False:
    print("     （产物与文献库不一致，先跑一次全流程）")
    _rq = urllib.request.Request(B + "/api/run",
                                 data=urllib.parse.urlencode({"stage": "all"}).encode(),
                                 method="POST")
    _rq.add_header("Content-Type", "application/x-www-form-urlencoded")
    try:
        urllib.request.urlopen(_rq, timeout=1800).read()
    except Exception as _e:
        print("     全流程执行失败：%s" % str(_e)[:80])
    st, j = get("/api/state")
ck("state.freshness 标记最新", j["freshness"]["needs_run"] is False, j["freshness"])
ck("state.check 含 missing_citations", "missing_citations" in j["check"])

st, pj = get("/api/papers")
ck("GET /api/papers 返回 200 且 6 篇", st == 200 and pj["count"] == 6, pj["count"])
ck("papers 每篇有标识符",
   all(p.get("local_pdf") for p in pj["papers"]), [p.get("local_pdf") for p in pj["papers"]])

st, pv = get("/api/privacy")
ck("GET /api/privacy 返回 200", st == 200 and "mode" in pv, pv)

_, ev = get("/api/evidence/detail?citation=1")
ck("证据详情返回文献与片段", bool(ev.get("paper")) and len(ev.get("evidence") or []) > 0,
   (ev.get("paper"), len(ev.get("evidence") or [])))
_, ev99 = get("/api/evidence/detail?citation=99")
ck("越界编号不崩（返回错误说明）", "error" in ev99 or ev99.get("evidence") == [], ev99.get("error"))
_, evx = get("/api/evidence/detail")
ck("缺参数不崩", isinstance(evx, dict))

# ══════════ 2. 隐私设置 ══════════
print()
print("=" * 70)
print("2. 隐私设置（四种入口）")
print("=" * 70)
ok, code = expect_http("/api/privacy?mode=cloud", "POST")
ck("POST /api/privacy?mode=cloud", ok, code)
_, pv = get("/api/privacy")
ck("设置后 mode=cloud 已确认", pv["mode"] == "cloud" and pv["acknowledged"] is True, pv)

ok, code = expect_http("/api/privacy?mode=local", "POST")
ck("POST /api/privacy?mode=local", ok, code)
_, pv = get("/api/privacy")
ck("切换为 local 生效", pv["mode"] == "local", pv["mode"])

# 非法 mode：接口应当拒绝（400），且不得把状态改成非法值
ok, code = expect_http("/api/privacy?mode=bogus", "POST", code=400)
ck("非法 mode 被拒绝（400）", ok and code == 400, code)
_, pv = get("/api/privacy")
ck("拒绝后模式仍合法", pv["mode"] in ("cloud", "local"), pv["mode"])
ok, code = expect_http("/api/privacy?mode=" + urllib.parse.quote("乱七八糟"), "POST", code=400)
ck("非 ASCII mode 被拒绝（400）", ok and code == 400, code)

# 磁盘上的隐私文件被写坏时，读取应归一化而不是卡住
import json as _json
_pf = os.path.join(PROJ, "data", "privacy.json")
_orig = open(_pf, encoding="utf-8").read() if os.path.isfile(_pf) else None
try:
    with open(_pf, "w", encoding="utf-8") as f:
        _json.dump({"mode": "corrupted", "acknowledged": True,
                    "acknowledged_at": "2026-01-01 00:00:00"}, f, ensure_ascii=False)
    _sys_mod = __import__("importlib")
    import wenmai.app as _app
    _sys_mod.reload(_app)
    ck("隐私文件被写坏时归一化为 cloud", _app._load_privacy()["mode"] == "cloud",
       _app._load_privacy()["mode"])
finally:
    if _orig is not None:
        with open(_pf, "w", encoding="utf-8") as f:
            f.write(_orig)
    else:
        os.remove(_pf)

# GET 跳转式设置：这是 fetch 走不通时的兜底路径。
# 它必须满足两点：① 设置确实生效；② **自动把用户送回应用**——
# 旧版只给一句"可以关闭本页"和一个手动链接，用户被丢在独立页面上，
# 实测被反馈为"很乱、像断开了"。文案随之改为「已切换为」，断言一并更新。
try:
    with urllib.request.urlopen(B + "/api/privacy/set?mode=cloud", timeout=30) as r:
        html = r.read().decode("utf-8", "ignore")
    ck("GET /api/privacy/set 返回确认页", r.status == 200 and "已切换为" in html, r.status)
    ck("确认页会自动返回应用（meta refresh）",
       "http-equiv='refresh'" in html and "mode_changed=" in html,
       "含 refresh" if "http-equiv='refresh'" in html else "缺 refresh")
    ck("确认页不再让用户自己找路回去",
       "正在自动返回应用" in html and "可以关闭本页" not in html)
except Exception as e:
    ck("GET /api/privacy/set 返回确认页", False, repr(e)[:70])

# 表单式保存（不跟随重定向，应 303）
code, hdrs = post_noredirect("/api/privacy/save", {"mode": "cloud"})
ck("POST /api/privacy/save 返回 303 重定向", code == 303, code)
ck("  └ Location 指向首页", (hdrs.get("Location") or hdrs.get("location")) == "/",
   hdrs.get("Location") or hdrs.get("location"))

# ══════════ 3. 页面与缓存头 ══════════
print()
print("=" * 70)
print("3. 页面路由与防缓存")
print("=" * 70)


def fetch(path):
    with urllib.request.urlopen(B + path, timeout=60) as r:
        return r.status, r.read().decode("utf-8", "ignore"), dict(r.headers)


st, html, hdr = fetch("/app")
ck("/app 返回主应用", st == 200 and "帮你把文献写成综述" in html)
ck("主应用含隐私弹窗（默认隐藏）", 'id="privacy-modal"' in html and 'class="hidden"' in html)
ck("主应用含防缓存 meta", "no-store" in html)
_, h_low = get_headers("/app")
ck("响应含 Cache-Control", "no-store" in (h_low.get("cache-control") or ""),
   h_low.get("cache-control"))
ck("响应含 Pragma", (h_low.get("pragma") or "") == "no-cache", h_low.get("pragma"))
ck("响应含版本标记 wm-version", "wm-version" in html)

st, html2, _ = fetch("/?skip=1")
ck("/?skip=1 直接进主应用", "帮你把文献写成综述" in html2)

st, html3, _ = fetch("/favicon.ico") if False else (0, "", {})
# favicon 是否提供
try:
    with urllib.request.urlopen(B + "/favicon.ico", timeout=10) as r:
        ck("favicon 提供（避免控制台 404）", True, r.status)
except urllib.error.HTTPError as e:
    ck("favicon 提供（避免控制台 404）", False, "HTTP %d" % e.code)

# ══════════ 4. 上传与删除 ══════════
print()
print("=" * 70)
print("4. 上传与删除")
print("=" * 70)
before = get("/api/papers")[1]["count"]

# 上传真实 PDF
pdf_path = os.path.join(PROJ, "data", "pdf", "006.pdf")
data = open(pdf_path, "rb").read()
b = uuid.uuid4().hex
body = ("--%s\r\n" % b).encode()
body += b'Content-Disposition: form-data; name="files"; filename="upload-test.pdf"\r\n'
body += b"Content-Type: application/pdf\r\n\r\n" + data + b"\r\n" + ("--%s--\r\n" % b).encode()
st, r = post_raw("/api/upload", body, "multipart/form-data; boundary=%s" % b)
ck("上传 PDF 成功", st == 200 and r["total"] == before + 1, (r.get("saved"), r.get("total")))

# 上传损坏的 PDF（扩展名对、内容错）
b2 = uuid.uuid4().hex
body2 = ("--%s\r\n" % b2).encode()
body2 += b'Content-Disposition: form-data; name="files"; filename="broken.pdf"\r\n'
body2 += b"Content-Type: application/pdf\r\n\r\nNOTAPDF" + b"\r\n" + ("--%s--\r\n" % b2).encode()
st, r2 = post_raw("/api/upload", body2, "multipart/form-data; boundary=%s" % b2)
ck("损坏文件被拒（skipped 有内容）", len(r2.get("skipped") or []) == 1, r2.get("skipped"))

# 上传不支持的格式
b3 = uuid.uuid4().hex
body3 = ("--%s\r\n" % b3).encode()
body3 += b'Content-Disposition: form-data; name="files"; filename="note.txt"\r\n'
body3 += b"Content-Type: text/plain\r\n\r\nhello" + b"\r\n" + ("--%s--\r\n" % b3).encode()
st, r3 = post_raw("/api/upload", body3, "multipart/form-data; boundary=%s" % b3)
ck("不支持的格式被拒", len(r3.get("skipped") or []) == 1, r3.get("skipped"))

# 上传真实 Word
try:
    from docx import Document as _D
    tmp = os.path.join(PROJ, ".dsh-tmp", "_up.docx")
    dd = _D()
    dd.add_heading("上传测试论文", 0)
    dd.add_heading("摘要", level=1)
    dd.add_paragraph("本文测试 Word 上传通道。" * 10)
    dd.save(tmp)
    wdata = open(tmp, "rb").read()
    b4 = uuid.uuid4().hex
    body4 = ("--%s\r\n" % b4).encode()
    body4 += b'Content-Disposition: form-data; name="files"; filename="upload-test.docx"\r\n'
    body4 += b"Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n"
    body4 += wdata + b"\r\n" + ("--%s--\r\n" % b4).encode()
    st, r4 = post_raw("/api/upload", body4, "multipart/form-data; boundary=%s" % b4)
    ck("上传 Word 成功", st == 200 and r4["total"] == before + 2, r4.get("saved"))
except Exception as e:
    ck("上传 Word 成功", False, repr(e)[:70])

# 上传后新鲜度应变为"需要重新生成"
st, j2 = get("/api/state")
ck("上传后 freshness.needs_run=True", j2["freshness"]["needs_run"] is True, j2["freshness"])

# 逐篇删除刚上传的
papers = get("/api/papers")[1]["papers"]
for tgt in [p for p in papers if "upload-test" in (p.get("title") or "")]:
    st, d = post("/api/delete", {"local_pdf": tgt["local_pdf"]})
    ck("删除 %s" % tgt["title"][:24], d.get("ok") is True, d)
ck("删除后回到 %d 篇" % before, get("/api/papers")[1]["count"] == before,
   get("/api/papers")[1]["count"])

# 删除不存在 / 越权
ok, code = expect_http("/api/delete", "POST", {"local_pdf": "pdf/nope.pdf"}, 404)
ck("删除不存在的记录返回 404", ok, code)
ok, code = expect_http("/api/delete", "POST", {"local_pdf": "../../../etc/passwd"}, 404)
ck("越权路径被拒（404）", ok, code)
ok, code = expect_http("/api/delete", "POST", {}, 400)
ck("缺参数返回 400", ok, code)

# ══════════ 5. 题录接口（错误路径） ══════════
print()
print("=" * 70)
print("5. 题录接口错误路径")
print("=" * 70)
ok, code = expect_http("/api/refs", "POST", {"text": "随便一段没有结构的文字"}, 400)
ck("无结构文本返回 400", ok, code)
ok, code = expect_http("/api/refs", "POST", {"text": ""}, 400)
ck("空题录返回 400", ok, code)
st, rr = post("/api/refs", {"text": "A1: 边界测试\nT1: 边界测试题录B\nJF: 测试刊\nYR: 2025\n"})
ck("正常题录导入成功", rr.get("added") == 1, rr)
refs_now = [p for p in get("/api/papers")[1]["papers"] if "边界测试题录B" in (p.get("title") or "")]
if refs_now:
    st, d = post("/api/delete", {"local_pdf": refs_now[0]["local_pdf"]})
    ck("题录可删除", d.get("ok") is True, d)

# ══════════ 6. 运行接口与正确性 ══════════
print()
print("=" * 70)
print("6. 运行接口：完整重跑与结果正确性")
print("=" * 70)
st, run = post("/api/run", {"stage": "all"}, timeout=900)
ck("POST /api/run stage=all", st == 200 and run.get("ok") is not False, run.get("log", [])[:2])
st, j3 = get("/api/state")
ck("重跑后 freshness 全部 current",
   all(j3["freshness"][k] == "current" for k in ("docs", "cards", "draft")), j3["freshness"])
ck("重跑后文献库仍 6 篇", j3["corpus"]["count"] == 6, j3["corpus"]["count"])
ck("重跑后自检有结果", "total" in j3["check"])
ck("重跑后埋错测试仍 4/4",
   j3.get("benchmark", {}).get("recall") == 1.0, j3.get("benchmark", {}).get("recall"))
ev = j3["check"].get("evidence_check", {})
ck("RAG 证据核对有支持率", ev.get("support_rate") is not None, ev.get("support_rate"))

# ══════════ 7. 下载接口 ══════════
print()
print("=" * 70)
print("7. 下载接口")
print("=" * 70)
for path, ext in [("/api/download/docx", ".docx"), ("/api/download/matrix", ".xlsx")]:
    try:
        with urllib.request.urlopen(B + path, timeout=120) as r:
            blob = r.read()
        ck("下载 %s（%d KB）" % (path, len(blob) // 1024), len(blob) > 1000, len(blob))
        magic = blob[:2] == b"PK"
        ck("  └ %s 是合法 Office 文件（PK 头）" % ext, magic, blob[:4])
    except Exception as e:
        ck("下载 %s" % path, False, repr(e)[:70])

# ══════════ 8. CLI 命令行 ══════════
print()
print("=" * 70)
print("8. CLI 命令行")
print("=" * 70)
CMDS = [["--help"], ["doctor"], ["status"], ["stop", "--port", "8999"], ["missing"]]
for args in CMDS:
    p = subprocess.run([sys.executable, "-X", "utf8", "-m", ENV.PKG] + args,
                       capture_output=True, text=True, cwd=PROJ, timeout=300,
                       encoding="utf-8", errors="ignore")
    ck("CLI %s 正常退出（code=%d）" % (" ".join(args), p.returncode), p.returncode == 0,
       (p.stderr or "")[:80])

p = subprocess.run([sys.executable, "-X", "utf8", "-m", ENV.PKG, "nosuchcmd"],
                   capture_output=True, text=True, cwd=PROJ, timeout=60,
                   encoding="utf-8", errors="ignore")
ck("未知命令报错退出（code!=0）", p.returncode != 0, p.returncode)

p = subprocess.run([sys.executable, "-X", "utf8", "-m", ENV.PKG, "refs", "--file", "不存在.txt"],
                   capture_output=True, text=True, cwd=PROJ, timeout=120,
                   encoding="utf-8", errors="ignore")
ck("refs 文件不存在时报错退出", p.returncode != 0, p.returncode)

# ══════════ 汇总 ══════════
print()
print("=" * 70)
print("第三批汇总：通过 %d，失败 %d" % (len(R["pass"]), len(R["fail"])))
print("=" * 70)
if R["fail"]:
    print("失败项：")
    for x in R["fail"]:
        print("  ✗ %s　→ %s" % (x["name"], x["detail"]))
with open(os.path.join(ENV.tmp_dir(), "testC_result.json"), "w", encoding="utf-8") as f:
    json.dump(R, f, ensure_ascii=False, indent=2)

# ── 收尾：关掉本次启动的服务 ──
try:
    _server_handle.stop()
except Exception:
    pass
